"""
Portal, BI e relatórios (Anexo I 2.7, 2.8, 2.9): isolamento, precisão dos
indicadores, filtros no servidor, paginação e consistência das exportações.

A massa é pequena e com datas fixas para que cada número esperado possa ser
conferido à mão — o cálculo de cada um está no comentário ao lado.
"""
import csv
import io
from datetime import date, datetime
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.cadastros.models import Area, Cliente, Equipamento, Setor, TecnologiaAnalise
from apps.coletas.models import Inspecao, MedicaoTecnica, Relatorio as RelatorioInspecao
from apps.laudos.models import Laudo
from apps.osp.models import OrdemServico

from . import indicadores as I
from .filtros import Filtros
from .reports import REPORTS

User = get_user_model()


def texto_do_pdf(conteudo: bytes) -> str:
    """Descomprime os fluxos de conteúdo do ReportLab (ASCII85 + Flate) para ler o texto."""
    import base64
    import re
    import zlib

    partes = []
    for bruto in re.findall(rb"stream\r?\n(.*?)endstream", conteudo, re.S):
        try:
            partes.append(zlib.decompress(base64.a85decode(b"<~" + bruto.strip(), adobe=True)).decode("latin-1"))
        except Exception:
            partes.append(bruto.decode("latin-1"))
    return "\n".join(partes)


def _norm(valor):
    """O Excel devolve 50.0 como 50: compara números como números."""
    try:
        return float(str(valor).replace(",", "."))
    except ValueError:
        return str(valor)


def quando(a, m, d, h=8):
    return timezone.make_aware(datetime(a, m, d, h))


class Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.a = Cliente.objects.create(nome="Indústria A", cnpj="11.111.111/0001-11")
        cls.b = Cliente.objects.create(nome="Indústria B", cnpj="22.222.222/0001-22")
        cls.vazio = Cliente.objects.create(nome="Indústria Vazia", cnpj="33.333.333/0001-33")
        cls.area_x = Area.objects.create(cliente=cls.a, nome="Utilidades")
        cls.area_y = Area.objects.create(cliente=cls.a, nome="Produção")
        cls.s1 = Setor.objects.create(area=cls.area_x, nome="Bombas")
        cls.s2 = Setor.objects.create(area=cls.area_y, nome="Prensas")
        sb = Setor.objects.create(area=Area.objects.create(cliente=cls.b, nome="Área B"), nome="Setor B")
        cls.a1 = Equipamento.objects.create(setor=cls.s1, tag="AA-01", nome="Bomba 1")
        cls.a2 = Equipamento.objects.create(setor=cls.s1, tag="AA-02", nome="Bomba 2")
        cls.a3 = Equipamento.objects.create(setor=cls.s2, tag="AA-03", nome="Prensa 3")
        cls.a4 = Equipamento.objects.create(setor=cls.s2, tag="AA-04", nome="Desativada", ativo=False)
        cls.b1 = Equipamento.objects.create(setor=sb, tag="BB-01", nome="Bomba da B")

        def u(email, perfil, cliente=None):
            return User.objects.create_user(email=email, nome=email.split("@")[0], perfil=perfil,
                                            cliente=cliente, password="SenhaSuperForte#2026!!")

        cls.gestor = u("gestor@t.com", "GESTOR")
        cls.tecnico = u("tecnico@t.com", "TECNICO")
        cls.pcm_a = u("pcm@a.com", "CLIENTE_PCM", cls.a)
        cls.manut_a = u("manut@a.com", "CLIENTE_MANUT", cls.a)
        cls.pcm_b = u("pcm@b.com", "CLIENTE_PCM", cls.b)
        cls.orfao = u("orfao@x.com", "CLIENTE_PCM", None)  # conta de cliente mal configurada

        # Medições sem disparar a geração automática de OSP (a massa de OSP é controlada abaixo).
        with patch("apps.osp.signals.processar_medicao"):
            cls.i1 = Inspecao.objects.create(cliente=cls.a, tecnico=cls.tecnico, data=date(2026, 1, 10),
                                             tipo_analise="SENSITIVA")
            cls.i2 = Inspecao.objects.create(cliente=cls.a, tecnico=cls.tecnico, data=date(2026, 3, 5),
                                             tipo_analise="SENSITIVA", status="CONCLUIDA")
            ib = Inspecao.objects.create(cliente=cls.b, tecnico=cls.tecnico, data=date(2026, 1, 15),
                                         tipo_analise="SENSITIVA")
            for insp, eq, crit in (
                (cls.i1, cls.a1, "CRITICO"), (cls.i1, cls.a2, "NORMAL"), (cls.i1, cls.a3, "ALERTA"),
                (cls.i1, cls.a4, "CRITICO"), (cls.i1, cls.a1, ""),
                (cls.i2, cls.a1, "CRITICO"), (cls.i2, cls.a3, "NORMAL"),
                (ib, cls.b1, "CRITICO"),
            ):
                MedicaoTecnica.objects.create(
                    inspecao=insp, equipamento=eq, tipo="SENSITIVA", subtipo="SENSITIVA",
                    ponto_medicao=f"P-{eq.tag}", grandeza="Parecer qualitativo", valor=None,
                    parametros={"criticidade": crit, "observacao": "teste"}, coletado_por=cls.tecnico,
                )

        def osp(eq, cliente, criado, **campos):
            o = OrdemServico.objects.create(cliente=cliente, equipamento=eq, titulo=f"OSP {eq.tag}",
                                            responsavel=cls.tecnico)
            OrdemServico.objects.filter(pk=o.pk).update(criado_em=criado, **campos)
            return o

        # O1: 24 h de reparo, dentro do prazo, custo real 100, ROI = 400 − 100 = 300.
        cls.o1 = osp(cls.a1, cls.a, quando(2026, 1, 10), status="FINALIZADA", finalizada_em=quando(2026, 1, 11),
                     sla_data=date(2026, 1, 20), custo_real=100, custo_estimado=150,
                     pred_mao_obra_valor=100, emerg_mao_obra_valor=400)
        # O2: mesmo equipamento 10 dias depois (MTBF), 72 h de reparo, fora do prazo, sem custo.
        cls.o2 = osp(cls.a1, cls.a, quando(2026, 1, 20), status="FINALIZADA", finalizada_em=quando(2026, 1, 23),
                     sla_data=date(2026, 1, 21))
        # O3: aberta e vencida.
        cls.o3 = osp(cls.a2, cls.a, quando(2026, 3, 1), status="ABERTA", sla_data=date(2026, 3, 2))
        # O4: cancelada — não entra em nada, nem o custo de 999.
        cls.o4 = osp(cls.a3, cls.a, quando(2026, 3, 2), status="CANCELADA", custo_real=999)
        cls.ob = osp(cls.b1, cls.b, quando(2026, 1, 15), status="ABERTA", custo_real=5000)

    def setUp(self):
        self.api = APIClient()

    def como(self, user):
        self.api.force_authenticate(user)
        return self.api


class PrecisaoDosIndicadores(Base):
    def test_operacional_do_cliente_a(self):
        d = I.operacional(self.gestor, Filtros(cliente=self.a.pk))
        self.assertEqual(d["total_inspecoes"], 2)
        self.assertEqual(d["inspecoes_abertas"], 1)
        self.assertEqual(d["medicoes_por_criticidade"], {"NORMAL": 2, "ALERTA": 1, "CRITICO": 3, "PENDENTE": 1})
        self.assertEqual(d["total_medicoes"], 7)
        # Só A1 conta: A4 é crítica mas está desativada. A1 teve 2 medições críticas.
        self.assertEqual([(e["equipamento__tag"], e["ocorrencias"]) for e in d["equipamentos_criticos"]], [("AA-01", 2)])
        # Disponibilidade (proxy): (3 ativos − 1 com crítica) / 3 = 66,7%.
        self.assertEqual(d["disponibilidade"]["percentual"], 66.7)
        # Falha recorrente: A1 tem 2 OSPs válidas.
        self.assertEqual(d["falhas_recorrentes"]["total_equipamentos"], 1)
        # SLA de atendimento: 2 finalizadas com prazo, 1 no prazo = 50%. 1 aberta vencida (O3).
        self.assertEqual(d["sla_atendimento"]["percentual"], 50.0)
        self.assertEqual(d["sla_atendimento"]["abertas_vencidas"], 1)
        self.assertEqual(d["osps_abertas"], 1)
        self.assertEqual(d["osps_total"], 3)  # O4 cancelada fora
        self.assertEqual(d["osps_por_status"], {"FINALIZADA": 2, "ABERTA": 1, "CANCELADA": 1})

    def test_executivo_do_cliente_a(self):
        d = I.executivo(self.gestor, Filtros(cliente=self.a.pk), hoje=date(2026, 3, 31))
        k, c = d["kpis"], d["custos"]
        self.assertEqual((k["osps_total"], k["osps_finalizadas"], k["osps_canceladas"]), (3, 2, 1))
        self.assertEqual(k["taxa_conclusao"], 67)    # 2/3
        self.assertEqual(k["mttr_horas"], 48.0)      # (24 + 72) / 2
        self.assertEqual(k["mtbf_dias"], 10.0)       # 10/01 → 20/01 no mesmo equipamento
        self.assertEqual(k["sla_cumprido_pct"], 50.0)
        self.assertEqual(c["real"], 100.0)           # 999 da cancelada fora; O2/O3 sem valor não viram 0
        self.assertEqual((c["osps_com_custo_real"], c["osps_sem_custo_real"]), (1, 2))
        self.assertEqual(c["estimado"], 150.0)
        self.assertEqual(c["retorno_investimento"], 300.0)
        self.assertEqual(c["osps_com_avaliacao"], 1)
        self.assertEqual([m["mes"] for m in d["evolucao"]],
                         ["10/2025", "11/2025", "12/2025", "01/2026", "02/2026", "03/2026"])
        jan = d["evolucao"][3]
        self.assertEqual((jan["criticas"], jan["osps"]), (2, 2))
        self.assertEqual(d["performance"], [
            {"cliente": "Indústria A", "osps": 3, "finalizadas": 2, "criticas": 3, "custo_real": 100.0},
        ])

    def test_periodo_restringe_volume_mas_nao_o_retrato_do_momento(self):
        jan = Filtros(cliente=self.a.pk, data_inicio=date(2026, 1, 1), data_fim=date(2026, 1, 31))
        d = I.operacional(self.gestor, jan)
        self.assertEqual(d["total_inspecoes"], 1)
        self.assertEqual(d["total_medicoes"], 5)
        self.assertEqual(d["osps_total"], 2)
        self.assertEqual(d["osps_abertas"], 1)  # O3 é de março, mas está aberta HOJE
        k = I.executivo(self.gestor, jan)["kpis"]
        self.assertEqual((k["osps_total"], k["taxa_conclusao"]), (2, 100))

    def test_ausencia_de_dado_e_none_e_nao_zero(self):
        f = Filtros(cliente=self.vazio.pk)
        d = I.executivo(self.gestor, f)
        self.assertEqual(d["kpis"]["osps_total"], 0)
        for chave in ("taxa_conclusao", "mttr_horas", "mtbf_dias", "sla_cumprido_pct"):
            self.assertIsNone(d["kpis"][chave], chave)
        for chave in ("real", "estimado", "retorno_investimento"):
            self.assertIsNone(d["custos"][chave], chave)
        self.assertEqual(d["performance"][0]["custo_real"], None)
        op = I.operacional(self.gestor, f)
        self.assertIsNone(op["disponibilidade"]["percentual"])
        self.assertIsNone(op["sla_atendimento"]["percentual"])

    def test_recorte_de_parque(self):
        d = I.operacional(self.gestor, Filtros(area=self.area_y.pk))
        # Área Y: A3 (ALERTA, NORMAL) e A4 (CRITICO, desativada).
        self.assertEqual(d["medicoes_por_criticidade"], {"NORMAL": 1, "ALERTA": 1, "CRITICO": 1, "PENDENTE": 0})
        self.assertEqual(d["equipamentos_criticos"], [])
        self.assertEqual(d["disponibilidade"]["equipamentos_ativos"], 1)
        self.assertEqual(d["total_inspecoes"], 2)  # as duas inspeções têm medição de A3
        self.assertEqual(d["osps_total"], 0)       # só O4 (cancelada) é da área Y

    def test_tela_e_exportacao_saem_da_mesma_conta(self):
        tela = self.como(self.gestor).get(f"/api/dashboard/?cliente={self.a.pk}").data
        csv_ = self.api.get(f"/api/relatorios/painel-operacional/?formato=csv&cliente={self.a.pk}")
        linhas = {r[0]: r[1] for r in csv.reader(io.StringIO(csv_.content.decode("utf-8-sig")), delimiter=";")}
        self.assertEqual(int(linhas["Medições no período"]), tela["total_medicoes"])
        self.assertEqual(float(linhas["Índice de disponibilidade (proxy)"]), tela["disponibilidade"]["percentual"])
        self.assertEqual(float(linhas["SLA de atendimento"]), tela["sla_atendimento"]["percentual"])
        self.assertEqual(int(linhas["OSPs abertas (hoje)"]), tela["osps_abertas"])


class Isolamento(Base):
    TEXTO_B = ("BB-01", "Indústria B", "Bomba da B")
    RELATORIOS_DO_PORTAL = [k for k, r in REPORTS.items() if not r["interno_only"]]

    def baixar(self, user, key, formato, extra=""):
        r = self.como(user).get(f"/api/relatorios/{key}/?formato={formato}&page_size=500{extra}")
        self.assertEqual(r.status_code, 200, (key, formato, getattr(r, "data", None)))
        if formato == "json":
            return str(r.data["linhas"])
        if formato == "csv":
            return r.content.decode("utf-8-sig")
        if formato == "xlsx":
            ws = load_workbook(io.BytesIO(r.content)).active
            return str([[c.value for c in row] for row in ws.iter_rows()])
        return texto_do_pdf(r.content)

    def test_cliente_nunca_ve_dado_de_outro_em_nenhum_formato(self):
        for key in self.RELATORIOS_DO_PORTAL:
            for formato in ("json", "csv", "xlsx", "pdf"):
                with self.subTest(key=key, formato=formato):
                    texto = self.baixar(self.pcm_a, key, formato)
                    for proibido in self.TEXTO_B:
                        self.assertNotIn(proibido, texto)
                    # e o parametro de cliente alheio não amplia: esvazia
                    texto = self.baixar(self.pcm_a, key, formato, f"&cliente={self.b.pk}")
                    for proibido in self.TEXTO_B:
                        self.assertNotIn(proibido, texto)

    def test_pdf_carrega_o_conteudo_do_escopo(self):
        """Garante que o teste negativo acima é significativo: o PDF tem texto legível."""
        texto = self.baixar(self.pcm_a, "equipamento", "pdf")
        self.assertIn("AA-01", texto)

    def test_conta_de_cliente_sem_empresa_nao_ve_nada(self):
        for key in self.RELATORIOS_DO_PORTAL:
            with self.subTest(key=key):
                r = self.como(self.orfao).get(f"/api/relatorios/{key}/?page_size=500")
                textos = str(r.data["linhas"])
                for proibido in ("AA-01", "BB-01", "Indústria A", "Indústria B"):
                    self.assertNotIn(proibido, textos)
        d = self.api.get("/api/dashboard/").data
        self.assertEqual((d["total_inspecoes"], d["total_medicoes"], d["osps_total"]), (0, 0, 0))
        p = self.api.get("/api/portal/visao-geral/").data
        self.assertEqual(p["equipamentos_atencao"], [])
        self.assertIsNone(p["indicadores"]["indice_disponibilidade"])

    def test_relatorios_e_painel_internos_bloqueados_ao_portal_em_todo_formato(self):
        for user in (self.pcm_a, self.manut_a):
            catalogo = {r["key"] for r in self.como(user).get("/api/relatorios/").data}
            self.assertFalse({"financeiro", "produtividade", "painel-executivo"} & catalogo)
            for key in ("financeiro", "produtividade", "painel-executivo"):
                for formato in ("json", "csv", "xlsx", "pdf"):
                    with self.subTest(user=user.perfil, key=key, formato=formato):
                        self.assertEqual(self.api.get(f"/api/relatorios/{key}/?formato={formato}").status_code, 403)
            self.assertEqual(self.api.get("/api/dashboard/executivo/").status_code, 403)

    def test_painel_e_portal_do_cliente_so_com_o_proprio_parque(self):
        d = self.como(self.pcm_b).get("/api/dashboard/").data
        self.assertEqual(d["total_medicoes"], 1)
        self.assertEqual([e["equipamento__tag"] for e in d["equipamentos_criticos"]], ["BB-01"])
        d = self.api.get(f"/api/dashboard/?cliente={self.a.pk}").data
        self.assertEqual((d["total_medicoes"], d["osps_total"]), (0, 0))
        p = self.como(self.pcm_a).get("/api/portal/visao-geral/").data
        self.assertEqual([e["equipamento__tag"] for e in p["equipamentos_atencao"]], ["AA-01"])
        self.assertEqual(p["indicadores"]["indice_disponibilidade"], 66.7)
        self.assertNotIn("BB-01", str(p))

    def test_listas_do_portal_respeitam_o_escopo(self):
        for rota in ("/api/osps/", "/api/inspecoes/", "/api/equipamentos/", "/api/laudos/"):
            with self.subTest(rota=rota):
                r = self.como(self.pcm_a).get(rota + "?page_size=500")
                self.assertEqual(r.status_code, 200)
                self.assertNotIn("BB-01", str(r.data))
                self.assertNotIn("Indústria B", str(r.data))

    def test_notificacoes_sao_so_do_proprio_usuario(self):
        """2.7.1.1.8 — o alerta de outra empresa (ou de outro usuário) não aparece nem por id."""
        from apps.notificacoes.models import Notificacao

        alheia = Notificacao.objects.create(destinatario=self.pcm_b, evento="NOVA_OSP", titulo="OSP da B",
                                            mensagem="BB-01")
        propria = Notificacao.objects.create(destinatario=self.pcm_a, evento="NOVA_OSP", titulo="OSP da A",
                                             mensagem="AA-01")
        api = self.como(self.pcm_a)
        lista = api.get("/api/notificacoes/?page_size=500").data
        ids = {n["id"] for n in lista["results"]}
        proprias = set(Notificacao.objects.filter(destinatario=self.pcm_a).values_list("id", flat=True))
        self.assertIn(propria.pk, ids)
        self.assertEqual(ids, proprias)  # as da massa (novas OSPs da A) + a criada aqui; nenhuma da B
        self.assertEqual(api.get(f"/api/notificacoes/{alheia.pk}/").status_code, 404)
        self.assertEqual(api.post(f"/api/notificacoes/{alheia.pk}/lida/").status_code, 404)
        self.assertEqual(api.get("/api/notificacoes/resumo/").data, {"nao_lidas": len(proprias)})
        api.post("/api/notificacoes/marcar-todas/")
        alheia.refresh_from_db()
        self.assertFalse(alheia.lida)

    def test_portal_conta_os_laudos_do_fluxo_novo(self):
        tec = TecnologiaAnalise.objects.create(nome="Vibração teste")
        rel = RelatorioInspecao.objects.create(cliente=self.a, tecnologia=tec, numero="RT-1",
                                               data_termino=date(2026, 3, 5))
        laudo = Laudo.objects.create(relatorio_origem=rel, titulo="Laudo do relatório", responsavel=self.tecnico)
        Laudo.objects.filter(pk=laudo.pk).update(status="EMITIDO", data_emissao=quando(2026, 3, 6))
        p = self.como(self.pcm_a).get("/api/portal/visao-geral/").data
        self.assertEqual(p["indicadores"]["laudos_disponiveis"], 1)
        self.assertEqual(self.como(self.pcm_b).get("/api/portal/visao-geral/").data["indicadores"]["laudos_disponiveis"], 0)


class FiltrosNoServidor(Base):
    def test_filtro_malformado_e_400_em_relatorio_e_painel(self):
        for qs in ("data_inicio=2026-13-01", "data_inicio=2026-02-01&data_fim=2026-01-01",
                   "cliente=abc", "setor=x"):
            with self.subTest(qs=qs):
                self.assertEqual(self.como(self.gestor).get(f"/api/relatorios/tecnico/?{qs}").status_code, 400)
                self.assertEqual(self.api.get(f"/api/relatorios/tecnico/?formato=csv&{qs}").status_code, 400)
                self.assertEqual(self.api.get(f"/api/dashboard/?{qs}").status_code, 400)
                self.assertEqual(self.api.get(f"/api/dashboard/executivo/?{qs}").status_code, 400)

    def test_relatorio_que_nao_aceita_recorte_de_parque_recusa_em_vez_de_ignorar(self):
        r = self.como(self.gestor).get(f"/api/relatorios/produtividade/?area={self.area_x.pk}")
        self.assertEqual(r.status_code, 400)

    def test_periodo_vale_em_todos_os_relatorios(self):
        """Antes, 'por equipamento' e 'produtividade' ignoravam as datas."""
        jan = "&data_inicio=2026-01-01&data_fim=2026-01-31"
        r = self.como(self.gestor).get(f"/api/relatorios/equipamento/?cliente={self.a.pk}{jan}").data
        linhas = {l[0]: l for l in r["linhas"]}
        self.assertEqual(linhas["AA-01"][3:], [2, 1, 2])  # 2 medições (1 crítica + 1 pendente), O1 e O2
        self.assertEqual(linhas["AA-03"][3:], [1, 0, 0])
        self.assertNotIn("AA-04", linhas)  # desativado
        r = self.api.get(f"/api/relatorios/produtividade/?cliente={self.a.pk}{jan}").data
        tecnico = next(l for l in r["linhas"] if l[0] == self.tecnico.nome)
        self.assertEqual(tecnico[2:], [1, 5, 0, 2])  # 1 inspeção, 5 medições, 0 laudos, O1 e O2 finalizadas em jan
        r = self.api.get(f"/api/relatorios/historico/?cliente={self.a.pk}{jan}").data
        self.assertEqual({l[0] for l in r["linhas"]}, {self.o1.numero, self.o2.numero})

    def test_relatorio_gerencial_pendente_fora_do_percentual(self):
        r = self.como(self.gestor).get(f"/api/relatorios/gerencial/?cliente={self.a.pk}").data
        # 7 medições, 1 pendente → 3 críticas / 6 classificadas = 50%.
        self.assertEqual(r["linhas"], [["Indústria A", 2, 7, 3, 1, "50%", 1]])
        r = self.api.get(f"/api/relatorios/gerencial/?cliente={self.vazio.pk}").data
        self.assertEqual(r["linhas"], [["Indústria Vazia", 0, 0, 0, 0, "—", 0]])

    def test_financeiro_nao_transforma_ausencia_em_zero_e_ignora_cancelada(self):
        r = self.como(self.gestor).get(f"/api/relatorios/financeiro/?cliente={self.a.pk}").data
        por_osp = {l[0]: l for l in r["linhas"]}
        self.assertNotIn(self.o4.numero, por_osp)
        self.assertEqual(por_osp[self.o1.numero][4:], ["R$ 150,00", "R$ 100,00", "R$ 300,00"])
        self.assertEqual(por_osp[self.o2.numero][4:], ["—", "—", "—"])
        self.assertEqual(por_osp["TOTAL"][1:], ["3 OSPs", "", "2 sem custo real", "R$ 150,00", "R$ 100,00", "R$ 300,00"])


class ExportacaoEPaginacao(Base):
    def test_json_csv_e_xlsx_trazem_as_mesmas_linhas(self):
        for key in REPORTS:
            with self.subTest(key=key):
                api = self.como(self.gestor)
                json_ = api.get(f"/api/relatorios/{key}/?page_size=500").data
                csv_ = api.get(f"/api/relatorios/{key}/?formato=csv")
                linhas_csv = list(csv.reader(io.StringIO(csv_.content.decode("utf-8-sig")), delimiter=";"))
                xlsx = load_workbook(io.BytesIO(api.get(f"/api/relatorios/{key}/?formato=xlsx").content)).active
                linhas_xlsx = [["" if c.value is None else str(c.value) for c in row] for row in xlsx.iter_rows()]
                esperado = [[str(c) for c in json_["colunas"]]] + [[str(c) for c in l] for l in json_["linhas"]]
                self.assertEqual(json_["total_linhas"], len(json_["linhas"]))
                self.assertEqual(linhas_csv, esperado)
                norm = lambda tabela: [[_norm(c) for c in l] for l in tabela]  # noqa: E731
                self.assertEqual(norm(linhas_xlsx), norm(esperado))
                pdf = api.get(f"/api/relatorios/{key}/?formato=pdf")
                self.assertEqual(pdf.status_code, 200)
                self.assertTrue(pdf.content.startswith(b"%PDF"))
                self.assertIn("attachment;", pdf["Content-Disposition"])

    def test_numero_vira_numero_no_excel(self):
        xlsx = load_workbook(io.BytesIO(self.como(self.gestor).get(
            f"/api/relatorios/gerencial/?formato=xlsx&cliente={self.a.pk}").content)).active
        self.assertEqual(xlsx["C2"].value, 7)
        self.assertIsInstance(xlsx["C2"].value, int)

    def test_texto_que_parece_formula_e_neutralizado(self):
        Equipamento.objects.filter(pk=self.a2.pk).update(tag='=HYPERLINK("http://x","clique")')
        api = self.como(self.gestor)
        csv_ = api.get(f"/api/relatorios/equipamento/?formato=csv&cliente={self.a.pk}").content.decode("utf-8-sig")
        self.assertIn("'=HYPERLINK", csv_)
        self.assertNotIn(';=HYPERLINK', csv_)
        self.assertFalse(any(l.startswith("=") for l in csv_.splitlines()))
        ws = load_workbook(io.BytesIO(api.get(
            f"/api/relatorios/equipamento/?formato=xlsx&cliente={self.a.pk}").content)).active
        tags = [row[0].value for row in ws.iter_rows(min_row=2)]
        self.assertIn("'=HYPERLINK(\"http://x\",\"clique\")", tags)

    def test_previa_paginada_e_arquivo_completo(self):
        api = self.como(self.gestor)
        r = api.get("/api/relatorios/tecnico/?page_size=3&page=2").data
        self.assertEqual((r["total_linhas"], r["paginas"], r["pagina"], len(r["linhas"])), (8, 3, 2, 3))
        r = api.get("/api/relatorios/tecnico/?page_size=99999").data
        self.assertEqual(r["page_size"], 500)
        csv_ = api.get("/api/relatorios/tecnico/?formato=csv&page_size=3").content.decode("utf-8-sig")
        self.assertEqual(len(csv_.strip().splitlines()), 1 + 8)  # o arquivo ignora a página

    def test_listas_respeitam_page_size_com_teto(self):
        for i in range(25):
            OrdemServico.objects.create(cliente=self.a, equipamento=self.a3, titulo=f"extra {i}")
        api = self.como(self.gestor)
        r = api.get("/api/osps/?page_size=30").data
        self.assertEqual(len(r["results"]), 30)  # antes: 20, em silêncio
        r = api.get("/api/osps/").data
        self.assertEqual(len(r["results"]), 20)
        from apps.core.pagination import PaginacaoPadrao
        self.assertEqual(PaginacaoPadrao.max_page_size, 500)
        r = api.get("/api/osps/?page_size=100000").data
        self.assertEqual(len(r["results"]), r["count"])  # 30 < 500: veio tudo, sem estourar

    def test_formato_invalido(self):
        self.assertEqual(self.como(self.gestor).get("/api/relatorios/tecnico/?formato=docx").status_code, 400)
